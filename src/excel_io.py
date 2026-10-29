from __future__ import annotations

import shutil
from dataclasses import dataclass
from datetime import datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.app_paths import data_dir
from src.maps_links import clean_maps_url

PROFILE_SHEET = "Perfil do Turista (Grupos)"
INVENTORY_SHEET = "Inventário Turístico"

INVENTORY_TEMPLATE_NAME = "Modelo_Inventario_Turistico.xlsx"

COMPACT_INVENTORY_HEADERS = (
    "ID",
    "Atração",
    "Abre",
    "Fecha",
    "Cap.",
    "Plog",
    "Cohen",
)


FULL_INVENTORY_HEADERS = (
    "ID",
    "Nome da Atração",
    "Categoria",
    "Horário de Abertura",
    "Horário de Fechamento",
    "Capacidade Máx. Grupo (Pessoas)",
    "Tempo Médio (horas)",
    "Peso Plog (Aventura)",
    "Peso Cohen (Independência)",
    "Descrição Breve",
    "Cidade",
    "Maps",
)

CIDADES_ROTEIRO = ("Botucatu", "Pardinho", "Bofete")


@dataclass(frozen=True)
class AxisQuestion:
    id: str
    pergunta: str
    significado_nota_1: str
    significado_nota_5: str


PLOG_QUESTIONS = [
    AxisQuestion(
        "P1",
        "Como o grupo prefere escolher as atividades?",
        "Atrações famosas, seguras e com infraestrutura",
        "Trilhas exóticas, aventuras e desafios",
    ),
    AxisQuestion(
        "P2",
        "Qual a preferência gastronômica do grupo?",
        "Restaurantes grandes ou comida familiar",
        "Comida de rua ou botecos locais",
    ),
    AxisQuestion(
        "P3",
        "Como o grupo reage a imprevistos na logística?",
        "Ansiedade alta / exige roteiro rígido",
        "Flexíveis, encaram imprevistos como aventura",
    ),
]

COHEN_QUESTIONS = [
    AxisQuestion(
        "C1",
        "Como foi o planejamento da viagem do grupo?",
        "Pacote fechado (pré-comprado)",
        "Passagem/hotel apenas; decide na hora",
    ),
    AxisQuestion(
        "C2",
        "Qual é a exigência de estrutura e conforto?",
        "Alto conforto (resorts, bons hotéis)",
        "Estrutura mínima (camping, hostels)",
    ),
    AxisQuestion(
        "C3",
        "Como o grupo pretende se deslocar?",
        "Ônibus fretado ou transfer agendado",
        "Transporte público ou a pé",
    ),
]

PERFIL_OPCOES = ["Casais", "Famílias", "Amigos", "Corporativo", "Estudantes", "Outro"]
RITMO_OPCOES = [
    "Relaxado (Poucas atividades)",
    "Moderado (Equilibrado)",
    "Intenso (Muitas atividades)",
]
RESTRICAO_OPCOES = ["Nenhuma", "Vegetariano", "Vegano", "Sem glúten", "Outra"]  # legado planilha modelo


def compute_group_scores(profile: TouristProfile) -> dict[str, float | str]:
    """Réplica das fórmulas da seção 4 da planilha Excel."""
    media_plog = profile.media_plog
    media_cohen = profile.media_cohen
    n = profile.numero_pessoas
    if n >= 15:
        fator = "Alto (Restringe atrações sem infra)"
    elif n >= 6:
        fator = "Médio"
    else:
        fator = "Baixo (Viagem ágil)"
    return {
        "media_plog": media_plog,
        "media_cohen": media_cohen,
        "fator_restritivo": fator,
    }


def parse_time_value(value: Any) -> time | None:
    if value is None or value == "":
        return None
    if isinstance(value, time):
        return value
    if isinstance(value, datetime):
        return value.time()
    text = str(value).strip()
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    return None


def format_time_value(value: time | None) -> str:
    if value is None:
        return "—"
    return value.strftime("%H:%M")


@dataclass
class TouristProfile:
    nome_grupo: str
    numero_pessoas: int
    duracao_dias: int
    perfil_predominante: str
    ritmo: str
    cidades: list[str]
    plog_scores: dict[str, int]
    cohen_scores: dict[str, int]

    @property
    def media_plog(self) -> float:
        vals = list(self.plog_scores.values())
        return sum(vals) / len(vals) if vals else 0.0

    @property
    def media_cohen(self) -> float:
        vals = list(self.cohen_scores.values())
        return sum(vals) / len(vals) if len(vals) else 0.0


@dataclass
class Attraction:
    id: str
    nome: str
    categoria: str
    horario_abertura: time | None
    horario_fechamento: time | None
    capacidade: int
    tempo_horas: float
    peso_plog: float
    peso_cohen: float
    descricao: str
    cidade: str = ""
    maps_url: str = ""

    @property
    def horario_funcionamento(self) -> str:
        if self.horario_abertura and self.horario_fechamento:
            return f"{format_time_value(self.horario_abertura)} às {format_time_value(self.horario_fechamento)}"
        return "Horário não informado"


def canonical_city_name(raw: str) -> str | None:
    text = (raw or "").strip().lower()
    if not text:
        return None
    for city in CIDADES_ROTEIRO:
        if city.lower() in text:
            return city
    return None


def parse_cidades_selecionadas(raw: Any) -> list[str]:
    text = str(raw or "").strip()
    if not text:
        return []
    selected: list[str] = []
    for part in text.replace(";", ",").split(","):
        canon = canonical_city_name(part.strip())
        if canon and canon not in selected:
            selected.append(canon)
    return selected


def serialize_cidades_selecionadas(cidades: list[str]) -> str:
    ordered = [c for c in CIDADES_ROTEIRO if c in cidades]
    return ", ".join(ordered)


def filter_inventory_by_cities(inventory: list[Attraction], cidades: list[str]) -> list[Attraction]:
    if not cidades:
        return []
    allowed = {c.lower() for c in cidades}
    filtered: list[Attraction] = []
    for item in inventory:
        canon = canonical_city_name(item.cidade)
        if canon and canon.lower() in allowed:
            filtered.append(item)
    return filtered


def default_workbook_path(project_root: Path | None = None) -> Path:
    root = project_root or data_dir().parent
    return root / "data" / "Projeto_Integrador_Turismo_Final.xlsx"


def ensure_workbook(path: Path, template_path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not path.exists():
        shutil.copy2(template_path, path)
    return path


def empty_tourist_profile() -> TouristProfile:
    """Perfil vazio (campos em branco até o usuário preencher)."""
    return TouristProfile(
        nome_grupo="",
        numero_pessoas=0,
        duracao_dias=0,
        perfil_predominante="",
        ritmo="",
        cidades=[],
        plog_scores={q.id: 0 for q in PLOG_QUESTIONS},
        cohen_scores={q.id: 0 for q in COHEN_QUESTIONS},
    )


def validate_profile_for_roteiro(profile: TouristProfile) -> list[str]:
    """Campos obrigatórios para gerar roteiro/PDF."""
    from src.form_constants import is_form_blank

    missing: list[str] = []
    if not profile.nome_grupo.strip():
        missing.append("Nome do grupo / contato")
    if profile.numero_pessoas < 1:
        missing.append("Número de pessoas")
    if profile.duracao_dias < 1:
        missing.append("Duração da viagem (dias)")
    if is_form_blank(profile.perfil_predominante):
        missing.append("Perfil predominante")
    if is_form_blank(profile.ritmo):
        missing.append("Ritmo de viagem")
    if not profile.cidades:
        missing.append("Cidades do roteiro (selecione ao menos uma)")
    for q in PLOG_QUESTIONS:
        if profile.plog_scores.get(q.id, 0) not in range(1, 6):
            missing.append(f"Plog {q.id} (escolha nota de 1 a 5)")
    for q in COHEN_QUESTIONS:
        if profile.cohen_scores.get(q.id, 0) not in range(1, 6):
            missing.append(f"Cohen {q.id} (escolha nota de 1 a 5)")
    return missing


def format_profile_validation_errors(missing: list[str]) -> str:
    if not missing:
        return ""
    bullets = "\n".join(f"• {item}" for item in missing)
    return f"Preencha todos os campos antes de gerar o roteiro:\n{bullets}"


def initialize_session_workbook(path: Path, template_path: Path) -> Path:
    """Copia a estrutura do modelo e zera perfil + inventário (sessão limpa)."""
    if not template_path.is_file():
        raise FileNotFoundError(f"Planilha modelo não encontrada: {template_path}")
    path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, path)
    wb = load_workbook(path)
    write_profile_to_sheet(wb[PROFILE_SHEET], empty_tourist_profile())
    if INVENTORY_SHEET in wb.sheetnames:
        del wb[INVENTORY_SHEET]
    inv_ws = wb.create_sheet(INVENTORY_SHEET)
    _write_full_inventory_sheet(inv_ws, [])
    wb.save(path)
    return path


def clear_profile_only(workbook_path: Path) -> Path:
    """Zera apenas o perfil na planilha atual; mantém o inventário importado."""
    if not workbook_path.is_file():
        raise FileNotFoundError(f"Planilha não encontrada: {workbook_path}")
    wb = load_workbook(workbook_path)
    if PROFILE_SHEET not in wb.sheetnames:
        raise ValueError(f"Aba ausente: {PROFILE_SHEET}")
    write_profile_to_sheet(wb[PROFILE_SHEET], empty_tourist_profile())
    wb.save(workbook_path)
    return workbook_path


def _cell_score_int(value: Any) -> int:
    if value is None or value == "":
        return 0
    if isinstance(value, str) and value.strip().startswith("="):
        return 0
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


def _read_profile_sheet(ws: Worksheet) -> TouristProfile:
    plog: dict[str, int] = {}
    cohen: dict[str, int] = {}

    for row in range(13, 16):
        qid = ws.cell(row, 1).value
        score = ws.cell(row, 3).value
        if qid:
            plog[str(qid)] = _cell_score_int(score)

    for row in range(19, 22):
        qid = ws.cell(row, 1).value
        score = ws.cell(row, 3).value
        if qid:
            cohen[str(qid)] = _cell_score_int(score)

    return TouristProfile(
        nome_grupo=str(ws["B4"].value or "").strip(),
        numero_pessoas=int(ws["B5"].value or 0),
        duracao_dias=int(ws["B6"].value or 0),
        perfil_predominante=str(ws["B7"].value or "").strip(),
        ritmo=str(ws["B8"].value or "").strip(),
        cidades=parse_cidades_selecionadas(ws["B9"].value),
        plog_scores=plog,
        cohen_scores=cohen,
    )


def write_profile_to_sheet(ws: Worksheet, profile: TouristProfile) -> None:
    ws["B4"] = profile.nome_grupo or None
    ws["B5"] = profile.numero_pessoas or None
    ws["B6"] = profile.duracao_dias or None
    ws["B7"] = profile.perfil_predominante or None
    ws["B8"] = profile.ritmo or None
    ws["B9"] = serialize_cidades_selecionadas(profile.cidades) or None

    row_map = {"P1": 13, "P2": 14, "P3": 15, "C1": 19, "C2": 20, "C3": 21}
    for qid, row in row_map.items():
        if qid.startswith("P"):
            val = profile.plog_scores.get(qid, 0)
        else:
            val = profile.cohen_scores.get(qid, 0)
        cell = ws.cell(row, 3)
        if val:
            cell.value = val
        else:
            cell.value = None
            cell.data_type = "n"


def save_profile(workbook_path: Path, profile: TouristProfile) -> None:
    wb = load_workbook(workbook_path)
    ws = wb[PROFILE_SHEET]
    write_profile_to_sheet(ws, profile)
    wb.save(workbook_path)


def read_profile(workbook_path: Path) -> TouristProfile:
    wb = load_workbook(workbook_path, data_only=False)
    return _read_profile_sheet(wb[PROFILE_SHEET])


def inventory_template_path(project_root: Path | None = None) -> Path:
    root = project_root or data_dir().parent
    return root / "data" / INVENTORY_TEMPLATE_NAME


def _normalize_header_cell(value: Any) -> str:
    return str(value or "").strip().lower().replace("ç", "c").replace("ã", "a")


def _inventory_format(header: tuple[Any, ...]) -> str:
    labels = [_normalize_header_cell(c) for c in header]
    if not labels or not labels[0]:
        return "unknown"
    if labels[0] == "id" and len(labels) >= 7:
        joined = " ".join(labels)
        if "abre" in labels and "fecha" in labels and "horario de abertura" not in joined:
            if "categoria" not in labels:
                return "compact"
    if any("abertura" in label for label in labels):
        return "full_hours"
    if labels[0] == "id" and len(labels) >= 9:
        return "full_legacy"
    return "unknown"


def _infer_categoria(nome: str) -> str:
    n = nome.lower()
    keys = ("feira", "cervejaria", "gastronomia", "restaurante", "comida", "café", "cafe")
    if any(k in n for k in keys):
        return "Gastronomia"
    if any(k in n for k in ("cachoeira", "mirante", "trilha", "natureza", "pedra")):
        return "Natureza"
    if any(k in n for k in ("museu", "pinacoteca", "catedral", "igreja", "cultura")):
        return "Cultura"
    return "Turismo"


def _infer_tempo_horas(categoria: str) -> float:
    if categoria == "Gastronomia":
        return 2.0
    if categoria == "Natureza":
        return 2.5
    return 1.5


def _full_inventory_has_city_column(header: tuple[Any, ...]) -> bool:
    labels = [_normalize_header_cell(c) for c in header if c]
    return any("cidade" in label for label in labels)


def _full_inventory_has_maps_column(header: tuple[Any, ...]) -> bool:
    labels = [_normalize_header_cell(c) for c in header if c]
    return any("maps" in label for label in labels)


def _column_index_by_header(header: tuple[Any, ...], needle: str) -> int | None:
    for idx, cell in enumerate(header):
        if needle in _normalize_header_cell(cell):
            return idx
    return None


def _normalize_full_import_row(row: tuple[Any, ...], header: tuple[Any, ...] | None = None) -> tuple[Any, ...]:
    """Converte linhas antigas (vegetariana) ou incompletas para o modelo atual."""
    cells = tuple(row)
    if header and _full_inventory_has_veg_column(header) and len(cells) >= 11:
        base = cells[:9] + (cells[10],)
        city = cells[11] if len(cells) > 11 else ""
        return base + (city,)
    if len(cells) >= len(FULL_INVENTORY_HEADERS):
        return cells[: len(FULL_INVENTORY_HEADERS)]
    padded = cells + ("",) * (len(FULL_INVENTORY_HEADERS) - len(cells))
    return padded[: len(FULL_INVENTORY_HEADERS)]


def _full_inventory_has_veg_column(header: tuple[Any, ...]) -> bool:
    labels = [_normalize_header_cell(c) for c in header if c]
    return any("veget" in label for label in labels)


def _parse_compact_row(
    row: tuple[Any, ...], header: tuple[Any, ...] | None = None
) -> Attraction | None:
    if not row or not row[0] or str(row[0]).upper() == "ID":
        return None
    try:
        nome = str(row[1] or "").strip()
        categoria = _infer_categoria(nome)
        descricao = nome
        plog_idx, cohen_idx = 5, 6
        maps_url = ""
        if header and _full_inventory_has_maps_column(header):
            maps_idx = _column_index_by_header(header, "maps")
            if maps_idx is not None and maps_idx == 5:
                plog_idx, cohen_idx = 6, 7
            if maps_idx is not None and len(row) > maps_idx:
                maps_url = clean_maps_url(row[maps_idx])
        return Attraction(
            id=str(row[0]).strip(),
            nome=nome,
            categoria=categoria,
            horario_abertura=parse_time_value(row[2]),
            horario_fechamento=parse_time_value(row[3]),
            capacidade=int(float(row[4] or 0)),
            tempo_horas=_infer_tempo_horas(categoria),
            peso_plog=float(row[plog_idx] or 0),
            peso_cohen=float(row[cohen_idx] or 0),
            descricao=descricao,
            cidade="",
            maps_url=maps_url,
        )
    except (TypeError, ValueError, IndexError):
        return None


def _write_full_inventory_sheet(ws: Worksheet, items: list[Attraction]) -> None:
    for col, header in enumerate(FULL_INVENTORY_HEADERS, start=1):
        ws.cell(1, col, header)
    for r, item in enumerate(items, start=2):
        ws.cell(r, 1, item.id)
        ws.cell(r, 2, item.nome)
        ws.cell(r, 3, item.categoria)
        ws.cell(r, 4, item.horario_abertura)
        ws.cell(r, 5, item.horario_fechamento)
        ws.cell(r, 6, item.capacidade)
        ws.cell(r, 7, item.tempo_horas)
        ws.cell(r, 8, item.peso_plog)
        ws.cell(r, 9, item.peso_cohen)
        ws.cell(r, 10, item.descricao)
        ws.cell(r, 11, item.cidade)
        ws.cell(r, 12, item.maps_url or None)


def _external_inventory_model_path() -> Path:
    return (
        Path.home()
        / "OneDrive"
        / "Desktop"
        / "PROJETO INTEGRADOR"
        / "Modelo_Inventario_Turistico.xlsx"
    )


def create_inventory_template(path: Path | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo

    target = path or inventory_template_path()
    target.parent.mkdir(parents=True, exist_ok=True)

    external = _external_inventory_model_path()
    if external.is_file():
        shutil.copy2(external, target)
        return target

    master = default_workbook_path(data_dir().parent)
    examples: list[tuple[Any, ...]] = []
    if master.exists():
        src = load_workbook(master, data_only=True)
        if INVENTORY_SHEET in src.sheetnames:
            ws_src = src[INVENTORY_SHEET]
            for row in ws_src.iter_rows(min_row=2, max_row=5, values_only=True):
                if row and row[0]:
                    examples.append(row)

    wb = Workbook()
    ws = wb.active
    ws.title = INVENTORY_SHEET

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")
    cap_header_fill = PatternFill("solid", fgColor="548235")
    score_header_fill = PatternFill("solid", fgColor="2F75B5")
    cap_fill = PatternFill("solid", fgColor="E2EFDA")
    score_fill = PatternFill("solid", fgColor="DDEBF7")

    for col, header in enumerate(FULL_INVENTORY_HEADERS, start=1):
        cell = ws.cell(1, col, header)
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if col == 6:
            cell.fill = cap_header_fill
        elif col in (8, 9):
            cell.fill = score_header_fill
        else:
            cell.fill = header_fill

    last_row = 1
    for row_idx, row in enumerate(examples, start=2):
        normalized = _normalize_full_import_row(row)
        for col_idx, value in enumerate(normalized[: len(FULL_INVENTORY_HEADERS)], start=1):
            ws.cell(row_idx, col_idx, value)
        last_row = row_idx

    if last_row == 1:
        last_row = 2
        ws.cell(2, 1, "ATR_01")

    for row in range(2, last_row + 1):
        ws.cell(row, 6).fill = cap_fill
        ws.cell(row, 8).fill = score_fill
        ws.cell(row, 9).fill = score_fill

    widths = [10, 34, 22, 14, 14, 16, 14, 14, 16, 48, 18, 42]
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A2"

    table = Table(
        displayName="InventarioTuristico",
        ref=f"A1:L{last_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    peso_validation = DataValidation(
        type="list",
        formula1='"1,2,3,4,5"',
        allow_blank=True,
    )
    peso_validation.error = "Escolha um valor de 1 a 5 na lista."
    peso_validation.errorTitle = "Peso inválido"
    ws.add_data_validation(peso_validation)
    peso_validation.add(f"H2:H500")
    peso_validation.add(f"I2:I500")

    city_validation = DataValidation(
        type="list",
        formula1='"Botucatu - SP,Pardinho - SP,Bofete - SP"',
        allow_blank=True,
    )
    ws.add_data_validation(city_validation)
    city_validation.add(f"K2:K500")

    help_ws = wb.create_sheet("Como preencher")
    help_ws["A1"] = f"Modelo de inventário — {INVENTORY_SHEET}"
    help_ws["A3"] = "Preencha todas as colunas conforme o cabeçalho da aba Inventário Turístico."
    help_ws["A4"] = "Horários: HH:MM. Peso Plog e Peso Cohen: escolha na lista (1 a 5)."
    help_ws["A5"] = "Coluna Cidade (K): Botucatu, Pardinho ou Bofete (ex.: Pardinho - SP)."
    help_ws["A6"] = "Coluna Maps (L): link do Google Maps (compartilhar → copiar link)."
    help_ws["A7"] = "Capacidade = número máximo de pessoas do grupo na atração."

    wb.save(target)
    return target


def _inventory_has_hours(header: tuple[Any, ...]) -> bool:
    return _inventory_format(header) == "full_hours"


def _parse_inventory_row(
    row: tuple[Any, ...],
    *,
    with_hours: bool,
    has_veg_column: bool = False,
    has_city_column: bool = False,
    has_maps_column: bool = False,
    header: tuple[Any, ...] | None = None,
) -> Attraction | None:
    if not row or not row[0] or str(row[0]).upper() == "ID":
        return None
    try:
        maps_url = ""
        if has_maps_column and header:
            maps_idx = _column_index_by_header(header, "maps")
            if maps_idx is not None and len(row) > maps_idx:
                maps_url = clean_maps_url(row[maps_idx])
        if with_hours:
            if has_veg_column and has_city_column:
                desc_idx, city_idx = 10, 11
            elif has_veg_column:
                desc_idx, city_idx = 10, None
            elif has_city_column:
                desc_idx, city_idx = 9, 10
            else:
                desc_idx, city_idx = 9, None
            cidade = (
                str(row[city_idx] or "").strip()
                if city_idx is not None and len(row) > city_idx
                else ""
            )
            return Attraction(
                id=str(row[0]).strip(),
                nome=str(row[1] or "").strip(),
                categoria=str(row[2] or "").strip(),
                horario_abertura=parse_time_value(row[3]),
                horario_fechamento=parse_time_value(row[4]),
                capacidade=int(float(row[5] or 0)),
                tempo_horas=float(row[6] or 0),
                peso_plog=float(row[7] or 0),
                peso_cohen=float(row[8] or 0),
                descricao=str(row[desc_idx] or "").strip() if len(row) > desc_idx else "",
                cidade=cidade,
                maps_url=maps_url,
            )
        return Attraction(
            id=str(row[0]).strip(),
            nome=str(row[1] or "").strip(),
            categoria=str(row[2] or "").strip(),
            horario_abertura=None,
            horario_fechamento=None,
            capacidade=int(float(row[3] or 0)),
            tempo_horas=float(row[4] or 0),
            peso_plog=float(row[5] or 0),
            peso_cohen=float(row[6] or 0),
            descricao=str(row[7] or "").strip(),
            cidade="",
            maps_url=maps_url,
        )
    except (TypeError, ValueError, IndexError):
        return None


def inventory_has_maps_column(workbook_path: Path) -> bool:
    wb = load_workbook(workbook_path, data_only=True)
    if INVENTORY_SHEET not in wb.sheetnames:
        return False
    header = next(wb[INVENTORY_SHEET].iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header:
        return False
    return _full_inventory_has_maps_column(header)


def inventory_maps_coverage(workbook_path: Path) -> tuple[int, int]:
    """Retorna (atrações com link Maps, total de atrações)."""
    items = read_inventory(workbook_path)
    if not items:
        return 0, 0
    with_maps = sum(1 for item in items if item.maps_url)
    return with_maps, len(items)


def read_inventory(workbook_path: Path) -> list[Attraction]:
    wb = load_workbook(workbook_path, data_only=True)
    ws = wb[INVENTORY_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = rows[0]
    fmt = _inventory_format(header)
    has_veg = _full_inventory_has_veg_column(header)
    has_city = _full_inventory_has_city_column(header)
    has_maps = _full_inventory_has_maps_column(header)
    items: list[Attraction] = []
    for row in rows[1:]:
        parsed: Attraction | None = None
        if fmt == "compact":
            parsed = _parse_compact_row(row, header)
        elif fmt == "full_hours":
            parsed = _parse_inventory_row(
                row,
                with_hours=True,
                has_veg_column=has_veg,
                has_city_column=has_city,
                has_maps_column=has_maps,
                header=header,
            )
        elif fmt == "full_legacy":
            parsed = _parse_inventory_row(
                row, with_hours=False, has_maps_column=has_maps, header=header
            )
        if parsed:
            items.append(parsed)
    return items


def read_inventory_from_uploaded(uploaded_path: Path) -> list[Attraction]:
    wb = load_workbook(uploaded_path, data_only=True)
    sheet = INVENTORY_SHEET if INVENTORY_SHEET in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = rows[0]
    fmt = _inventory_format(header)
    has_veg = _full_inventory_has_veg_column(header)
    has_city = _full_inventory_has_city_column(header)
    has_maps = _full_inventory_has_maps_column(header)
    items: list[Attraction] = []
    for row in rows[1:]:
        parsed: Attraction | None = None
        if fmt == "compact":
            parsed = _parse_compact_row(row, header)
        elif fmt == "full_hours":
            parsed = _parse_inventory_row(
                row,
                with_hours=True,
                has_veg_column=has_veg,
                has_city_column=has_city,
                has_maps_column=has_maps,
                header=header,
            )
        elif fmt == "full_legacy":
            parsed = _parse_inventory_row(
                row, with_hours=False, has_maps_column=has_maps, header=header
            )
        if parsed:
            items.append(parsed)
    return items


def import_inventory_sheet(workbook_path: Path, uploaded_path: Path) -> int:
    """Import inventory; compact model is normalized to full sheet in the project workbook."""
    items = read_inventory_from_uploaded(uploaded_path)
    if not items:
        raise ValueError("Nenhuma atração válida encontrada. Use o modelo Modelo_Inventario_Turistico.xlsx.")

    dst = load_workbook(workbook_path)
    if INVENTORY_SHEET in dst.sheetnames:
        del dst[INVENTORY_SHEET]
    dst_ws = dst.create_sheet(INVENTORY_SHEET)
    _write_full_inventory_sheet(dst_ws, items)
    dst.save(workbook_path)
    return len(items)


def profile_and_inventory_as_text(workbook_path: Path) -> str:
    profile = read_profile(workbook_path)
    inventory = read_inventory(workbook_path)
    lines = [
        "=== Perfil do Turista (Grupos) ===",
        f"Nome: {profile.nome_grupo}",
        f"Pessoas: {profile.numero_pessoas}",
        f"Duração (dias): {profile.duracao_dias}",
        f"Perfil: {profile.perfil_predominante}",
        f"Ritmo: {profile.ritmo}",
        f"Cidades: {', '.join(profile.cidades) if profile.cidades else '—'}",
        f"Média Plog: {profile.media_plog:.2f}",
        f"Média Cohen: {profile.media_cohen:.2f}",
        "",
        "=== Inventário Turístico ===",
    ]
    for a in inventory:
        lines.append(
            f"{a.id} | {a.nome} | {a.horario_funcionamento} | cap={a.capacidade} | "
            f"plog={a.peso_plog} | cohen={a.peso_cohen} | tempo={a.tempo_horas}h | "
            f"cidade={a.cidade or '—'} | maps={a.maps_url or '—'} | {a.descricao}"
        )
    return "\n".join(lines)
